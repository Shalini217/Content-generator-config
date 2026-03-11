from openpyxl import load_workbook
import json

# Read the Excel file
excel_file = r"C:\Users\DELL\Downloads\Content Types.xlsx"

try:
    # Load workbook
    wb = load_workbook(excel_file, read_only=True, data_only=True)
    ws = wb.active

    print(f"Reading Excel file: {excel_file}")
    print(f"Active sheet: {ws.title}")
    print("="*60)

    # Get headers from first row
    headers = []
    for cell in ws[1]:
        if cell.value:
            headers.append(str(cell.value))

    print(f"Headers found: {headers}")
    print("="*60)

    # Find column indices
    content_type_col = None
    product_field_col = None

    for idx, header in enumerate(headers):
        header_lower = header.lower().replace(' ', '').replace('_', '')
        if header_lower in ['contenttype', 'content']:
            content_type_col = idx
            print(f"Content Type column: {idx} ({header})")
        elif 'productfield' in header_lower or 'product' in header_lower:
            product_field_col = idx
            print(f"Product Field column: {idx} ({header})")

    if content_type_col is None:
        content_type_col = 0  # Default to first column
        print(f"Using first column (0) for Content Type")

    if product_field_col is None:
        product_field_col = 1  # Default to second column
        print(f"Using second column (1) for Product Field")

    print("="*60)

    # Extract data
    content_types_data = {}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row and len(row) > max(content_type_col, product_field_col):
            content_type = row[content_type_col]
            product_field = row[product_field_col]

            if content_type and product_field:
                content_type_str = str(content_type).strip()
                product_field_str = str(product_field).strip()

                if content_type_str and product_field_str:
                    content_types_data[content_type_str] = {
                        "productFieldName": product_field_str
                    }
                    print(f"Row {row_idx}: {content_type_str} -> {product_field_str}")

    wb.close()

    print("\n" + "="*60)
    print("EXTRACTED DATA (JSON):")
    print("="*60)
    print(json.dumps(content_types_data, indent=4))

    print("\n" + "="*60)
    print("JAVASCRIPT FORMAT FOR HTML FILE:")
    print("="*60)
    print("const CONTENT_TYPES_DATA = {")
    items = list(content_types_data.items())
    for idx, (ct_name, ct_data) in enumerate(items):
        comma = "," if idx < len(items) - 1 else ""
        print(f'    "{ct_name}": {{')
        print(f'        "productFieldName": "{ct_data["productFieldName"]}"')
        print(f'    }}{comma}')
    print("};")

    print("\n" + "="*60)
    print(f"✓ Successfully extracted {len(content_types_data)} content types")
    print("="*60)

    # Save to file for easy copying
    output_file = r"C:\Users\DELL\content_types_data.txt"
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write("const CONTENT_TYPES_DATA = {\n")
        items = list(content_types_data.items())
        for idx, (ct_name, ct_data) in enumerate(items):
            comma = "," if idx < len(items) - 1 else ""
            f.write(f'    "{ct_name}": {{\n')
            f.write(f'        "productFieldName": "{ct_data["productFieldName"]}"\n')
            f.write(f'    }}{comma}\n')
        f.write("};\n")

    print(f"\n✓ Output saved to: {output_file}")
    print("  Copy this content into your HTML file!")

except FileNotFoundError:
    print(f"❌ Error: File not found at {excel_file}")
    print("Please make sure 'Content Types.xlsx' exists in your Downloads folder")
except Exception as e:
    print(f"❌ Error: {e}")
    import traceback
    traceback.print_exc()
